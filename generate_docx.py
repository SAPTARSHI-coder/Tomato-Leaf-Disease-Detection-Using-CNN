import subprocess
import sys
import re
import os

try:
    import docx
except ImportError:
    print("Installing python-docx...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_word_document():
    doc = Document()
    
    # Configure base font to Times New Roman
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    md_file_path = os.path.join(os.path.dirname(__file__), 'Tomato_Leaf_Disease_Report.md')
    with open(md_file_path, 'r', encoding='utf-8') as f:
        text = f.read()

    # Split markdown into paragraphs by double newline
    paragraphs = text.split('\n\n')

    for p_text in paragraphs:
        p_text = p_text.strip()
        if not p_text:
            continue
        
        # Headings
        if p_text.startswith('# '):
            h = doc.add_heading(p_text[2:].strip(), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p_text.startswith('## '):
            doc.add_heading(p_text[3:].strip(), level=2)
        elif p_text.startswith('### '):
            doc.add_heading(p_text[4:].strip(), level=3)
            
        # Bullet list
        elif p_text.startswith('- ') or p_text.startswith('* '):
            lines = p_text.split('\n')
            for line in lines:
                line = line.strip()
                if line.startswith('- ') or line.startswith('* '):
                    p = doc.add_paragraph(style='List Bullet')
                    line_content = line[2:].strip()
                else:
                    # Append loosely to bullet logic
                    p = doc.add_paragraph()
                    line_content = line
                
                # Split and parse inline bold and italic markings
                parts = re.split(r'(\*\*.*?\*\*|\_.*?\_)', line_content)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        p.add_run(part[2:-2]).bold = True
                    elif part.startswith('_') and part.endswith('_'):
                        p.add_run(part[1:-1]).italic = True
                    else:
                        p.add_run(part)
        else:
            # Regular Text
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            # Extract Image/Table Placeholders explicitly and Bold/Italicize them
            parts = re.split(r'(\*\*.*?\*\*|\_.*?\_|\[Figure.*?\]|\[Table.*?\])', p_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                elif part.startswith('_') and part.endswith('_'):
                    p.add_run(part[1:-1]).italic = True
                elif part.startswith('[Figure') or part.startswith('[Table'):
                    run = p.add_run(part)
                    run.bold = True
                    run.italic = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.add_run(part)

    output_path = os.path.join(os.path.dirname(__file__), 'Tomato_Leaf_Disease_Report.docx')
    doc.save(output_path)
    print(f"Successfully generated DOCX format at: {output_path}")

if __name__ == "__main__":
    create_word_document()
