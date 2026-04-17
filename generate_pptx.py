import subprocess
import sys
import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("Installing python-pptx...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN

def add_slide(prs, title, content_bullets):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    title_shape.text = title
    
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.clear() # clear default paragraphs
    
    for i, bullet in enumerate(content_bullets):
        p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(22)
        p.level = 0 if not bullet.startswith("  -") else 1
        if p.level == 1:
            p.font.size = Pt(20)

def create_presentation():
    prs = Presentation()
    
    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Tomato Leaf Disease Detection Using CNN"
    subtitle.text = "Final Year B.Tech Project Presentation\n\nPresented by: [Student Name]"
    
    # Slide 2: Introduction
    add_slide(prs, "Introduction", [
        "Agriculture drives global food security and economic stability.",
        "Tomato crops are highly susceptible to bacterial, viral, and fungal diseases.",
        "Traditional manual detection is slow, subjective, and labor-intensive.",
        "Late diagnosis results in heavy yield reduction and financial losses.",
        "An automated, computer vision-based disease diagnosis system is highly needed."
    ])

    # Slide 3: Impact of Tomato Leaf Diseases
    add_slide(prs, "Impact of Tomato Leaf Diseases", [
        "Pathogens cause lesions, necrosis, and massive defoliation.",
        "Diseases spread rapidly across the entire agricultural field.",
        "Severely stunts plant growth and reduces fruit quality.",
        "High degree of visual similarity between different diseases at early stages.",
        "Expert agronomists are rare in rural farming locations."
    ])

    # Slide 4: Role of AI & Proposed Solution
    add_slide(prs, "Proposed AI Solution", [
        "Deploy Deep Learning and Convolutional Neural Networks (CNN) for image analysis.",
        "CNNs automatically extract complex visual patterns directly from pixels.",
        "Capable of distinguishing between healthy and diseased leaf classes rapidly.",
        "Scalable technology capable of serving farmers directly via mobile apps.",
        "Guarantees standard, objective, and non-fatiguing real-time disease diagnosis."
    ])

    # Slide 5: Literature Review & Gap
    add_slide(prs, "Literature Review & Research Gap", [
        "Traditional ML (SVM, K-Means) required manual, error-prone feature extraction.",
        "Pre-trained DL models (ResNet, VGG16) often carry massive computational weight.",
        "Many models fail to generalize outside controlled laboratory conditions.",
        "Research Gap: The need for a lightweight, customized CNN model balancing high inference speed and validation accuracy.",
        "Focus on robustness against varied lighting and mobile camera artifacts."
    ])

    # Slide 6: System Architecture Pipeline
    add_slide(prs, "System Architecture Pipeline", [
        "1. Image Acquisition: Capture leaf imagery via mobile/camera.",
        "2. Preprocessing: Resizing, normalization, and heavy data augmentation.",
        "3. Feature Extraction: Convolutional and Max Pooling sequential layers.",
        "4. Feature Synthesis: Fully Connected (Dense) neural network.",
        "5. Classification: Softmax activation yielding disease class probabilities."
    ])

    # Slide 7: Dataset & Preprocessing
    add_slide(prs, "Dataset & Preprocessing", [
        "Source: PlantVillage distinct Tomato subset combined with local datasets.",
        "Categorized into 10 classes (1 healthy, 9 specific diseases like Early Blight).",
        "Resizing: Images standardized (e.g., 224x224) to maintain tensor shape.",
        "Normalization: Pixel scales reduced from 0-255 to [0, 1] for optimization.",
        "Augmentation: Applied random rotation, flipping, and shearing to prevent overfitting."
    ])

    # Slide 8: CNN Model Architecture
    add_slide(prs, "CNN Model Architecture", [
        "Convolutional Layers: Extract fundamental patterns and hierarchical lesions.",
        "ReLU Activation: Introduces non-linearity to process complex inter-relations.",
        "Max Pooling Layers: Downsamples structural dimensions to lower processing limits.",
        "Flatten & Dense Layers: Synthesizes arrayed spatial vectors into class weights.",
        "Softmax Function: Converts final vector metrics into percentages (confidence)."
    ])

    # Slide 9: Model Training
    add_slide(prs, "Model Training Setup", [
        "Loss Function: Categorical Cross-Entropy (for multi-class problems).",
        "Optimizer: Adam (Adaptive Moment Estimation) for dynamic learning rate adjustments.",
        "Epochs: 50 iterative cycles with a batch size of 32 for smooth gradients.",
        "Evaluation Metrics: Tracking validation Accuracy, Precision, Recall, and F1-score.",
        "Deployment Formats: Exported into lightweight HDF5 or TFLite structures."
    ])

    # Slide 10: Results - Accuracy & Loss
    add_slide(prs, "Results: Accuracy and Loss", [
        "Achieved extremely high classification accuracy between training and testing splits.",
        "Validation Accuracy securely trailed Training Accuracy indicating minimal overfitting.",
        "Categorical Cross-Entropy Loss showed an aggressive smooth downward trajectory.",
        "NOTE: Ensure to add Accuracy vs Epoch Graphs here.",
        "NOTE: Ensure to add Model Loss Analysis Graphs here."
    ])

    # Slide 11: Confusion Matrix & Case Studies
    add_slide(prs, "Confusion Matrix & Case Studies", [
        "Confusion Matrix reveals highly concentrated True Positives on the true diagonal.",
        "Healthy Leaves: Identified perfectly via uniform green continuity.",
        "Early Blight: Effectively localized by concentric ring pattern detection.",
        "Mosaic Virus: Distinguished by uneven light/dark leaf vein manifestations.",
        "Minimal errors isolated between heavily similar necrotic categories."
    ])

    # Slide 12: Advantages & Limitations
    add_slide(prs, "Advantages & Limitations", [
        "Advantages:",
        "  - Zero reliance on manual feature extraction.",
        "  - Democratization of expensive laboratory expertise to localized farmers.",
        "Limitations:",
        "  - Solely relies on leaf symptoms; ignores stems and soil factors.",
        "  - Can be occasionally tricked by intense direct sun glare and shadow."
    ])

    # Slide 13: Conclusion & Future Scope
    add_slide(prs, "Conclusion & Future Scope", [
        "Conclusion: CNNs represent a robust solution for automated agronomics, replacing human subjectivity with mathematical certainty.",
        "Future Scope:",
        "  - Scale the system using Object Detection algorithms (e.g., YOLOv8).",
        "  - Integration with continuous drone video-feed evaluation.",
        "  - Deployment on cross-platform mobile apps for absolute accessibility."
    ])

    # Slide 14: Q&A
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Thank You!"
    subtitle = slide.placeholders[1]
    subtitle.text = "Any Questions?\n\nContact: [Your Name / Email]"

    output_path = os.path.join(os.path.dirname(__file__), 'Tomato_Leaf_Disease_Presentation.pptx')
    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == '__main__':
    create_presentation()
